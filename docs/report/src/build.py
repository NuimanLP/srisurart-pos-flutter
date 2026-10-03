"""Build the progress-report .docx from chapters/*.md on top of the faculty .dotx template.

Usage: python build.py <out.docx>
Reads: tpl source (.dotx), chapters/ch1.md ch2.md ch3.md ch4a.md ch4b.md ch5.md, diagrams/<slug>.png,
       cover.json. Repo images are resolved against REPO.
"""
import copy, json, os, re, sys, zipfile
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(HERE, "..", "..", ".."))
DOTX = os.environ.get("REPORT_DOTX", os.path.join(HERE, "project-report-template.dotx"))
CHAPTER_FILES = [["ch1.md"], ["ch2.md"], ["ch3.md"], ["ch4a.md", "ch4b.md"], ["ch5.md"]]

# ---------------------------------------------------------------- template -> docx
def dotx_to_docx(src, dst):
    with zipfile.ZipFile(src) as zi, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zo:
        for item in zi.infolist():
            data = zi.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(b"wordprocessingml.template.main+xml",
                                    b"wordprocessingml.document.main+xml")
            zo.writestr(item, data)

# ---------------------------------------------------------------- markdown parsing
def split_chapter(text):
    """Return (body_lines, diagrams{slug:spec}, abbrevs{abbr:full}, refs{n:text})."""
    diagrams = {}
    def grab(m):
        diagrams[m.group(1)] = m.group(2).strip()
        return ""
    text = re.sub(r"```diagram\s+slug=([\w\-]+)\s*\n(.*?)```", grab, text, flags=re.S)
    parts = re.split(r"^## __(ABBREVIATIONS|REFERENCES|FACTS)__\s*$", text, flags=re.M)
    body = parts[0]
    abbrevs, refs = {}, {}
    for i in range(1, len(parts), 2):
        kind, chunk = parts[i], parts[i + 1]
        for line in chunk.splitlines():
            line = line.strip()
            if kind == "ABBREVIATIONS" and "=" in line:
                a, f = line.split("=", 1)
                a = a.strip().strip("-* ").strip("`")
                if a:
                    abbrevs.setdefault(a, f.strip())
            elif kind == "REFERENCES":
                m = re.match(r"\[(\d+)\]\s*(.+)", line.lstrip("-* "))
                if m:
                    refs[int(m.group(1))] = m.group(2).strip()
    return body.splitlines(), diagrams, abbrevs, refs


def ref_key(text):
    m = re.search(r"https?://[^\s,\]>]+", text)
    return (m.group(0).rstrip(".").rstrip("/").lower() if m else re.sub(r"\W+", "", text.lower())[:80])

# ---------------------------------------------------------------- xml helpers
def set_cs_size(run, half_points):
    rpr = run._r.get_or_add_rPr()
    for tag in ("w:sz", "w:szCs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = OxmlElement(tag); rpr.append(el)
        el.set(qn("w:val"), str(half_points))


def field_runs(paragraph, instr, cached, size=None):
    def r(child):
        run = paragraph.add_run()
        run._r.append(child)
        if size: set_cs_size(run, size)
        return run
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin"); r(b)
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr; r(it)
    s = OxmlElement("w:fldChar"); s.set(qn("w:fldCharType"), "separate"); r(s)
    run = paragraph.add_run(cached)
    if size: set_cs_size(run, size)
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end"); r(e)


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def table_borders(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "4"); el.set(qn("w:space"), "0"); el.set(qn("w:color"), "808080")
        borders.append(el)
    tblPr.append(borders)
    jc = OxmlElement("w:jc"); jc.set(qn("w:val"), "center"); tblPr.append(jc)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    h = OxmlElement("w:tblHeader"); h.set(qn("w:val"), "true"); trPr.append(h)

def thai_just(p):
    pPr = p._p.get_or_add_pPr(); jc = pPr.find(qn("w:jc"))
    if jc is None:
        jc = OxmlElement("w:jc"); pPr.append(jc)
    jc.set(qn("w:val"), "thaiDistribute")

# ---------------------------------------------------------------- builder
class Builder:
    def __init__(self, doc):
        self.doc = doc
        self.ch = 0
        self.fig_no = self.tab_no = 0
        self.labels = {}           # 'fig:slug' -> '3.2'
        self.ref_order = []        # global list of ref texts
        self.ref_index = {}        # key -> global n
        self.local_map = {}        # local n -> global n (current chapter file)
        self.missing = []

    # -- inline text with `code`, **bold**, [n] citations, {fig:x} refs
    def inline(self, p, text, size=None, bold=False):
        tokens = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
        for tok in tokens:
            if not tok:
                continue
            if tok.startswith("`") and tok.endswith("`") and len(tok) > 1:
                run = p.add_run(re.sub(r"([/_.\-])", "\\1\u200b", tok[1:-1]))
                run.font.name = "Consolas"
                rpr = run._r.get_or_add_rPr(); rf = rpr.find(qn("w:rFonts"))
                rf.set(qn("w:cs"), "Consolas")
                set_cs_size(run, (size or 32) - 8)
                run.font.color.rgb = RGBColor(0x1F, 0x3B, 0x6E)
                if bold: run.bold = True
                continue
            b = bold
            if tok.startswith("**") and tok.endswith("**"):
                tok, b = tok[2:-2], True
            tok = self.rewrite_refs(tok)
            # yellow highlight for [ ... ] fill-in placeholders written as [[...]]
            for seg in re.split(r"(\[\[[^\]]+\]\])", tok):
                if not seg:
                    continue
                if seg.startswith("[["):
                    run = p.add_run(seg[2:-2]); run.font.highlight_color = WD_COLOR_INDEX.YELLOW
                else:
                    run = p.add_run(seg)
                if re.search(r"[฀-๿]", seg):
                    rpr = run._r.get_or_add_rPr()
                    rpr.append(OxmlElement("w:cs"))
                    lang = OxmlElement("w:lang"); lang.set(qn("w:bidi"), "th-TH"); rpr.append(lang)
                if b: run.bold = True; run._r.get_or_add_rPr().append(OxmlElement("w:bCs"))
                if size: set_cs_size(run, size)

    def rewrite_refs(self, tok):
        def cite(m):
            nums = []
            for part in re.split(r"\s*,\s*", m.group(1)):
                rng = re.match(r"(\d+)\s*[–-]\s*(\d+)$", part)
                if rng:
                    nums += list(range(int(rng.group(1)), int(rng.group(2)) + 1))
                elif part.isdigit():
                    nums.append(int(part))
            out = []
            for n in nums:
                g = self.local_map.get(n)
                if g is None:
                    self.missing.append(f"ch{self.ch}: citation [{n}] has no reference")
                    continue
                out.append(g)
            return "".join(f"[{g}]" for g in sorted(set(out))) if out else ""
        tok = re.sub(r"\[(\d+(?:\s*[,–-]\s*\d+)*)\]", cite, tok)
        def lab(m):
            key = m.group(1)
            if key not in self.labels:
                self.missing.append(f"ch{self.ch}: unresolved label {key}")
                return "?"
            return self.labels[key]
        return re.sub(r"\{((?:fig|tab):[\w\-]+)\}", lab, tok)

    def para(self, text, style="Body Text", indent_tab=True):
        p = self.doc.add_paragraph(style=style)
        p.paragraph_format.space_before = Pt(3); p.paragraph_format.space_after = Pt(3)
        if indent_tab:
            p.add_run().add_tab()
        self.inline(p, text)
        thai_just(p)
        return p

    def list_item(self, marker, text):
        p = self.doc.add_paragraph(style="Body Text")
        pf = p.paragraph_format
        pf.left_indent = Cm(1.9); pf.first_line_indent = Cm(-0.7)
        pf.space_before = Pt(0); pf.space_after = Pt(0)
        p.add_run(marker + "\t")
        pf.tab_stops.add_tab_stop(Cm(1.9))
        self.inline(p, text)
        thai_just(p)

    def heading(self, level, text):
        p = self.doc.add_paragraph(style=f"Heading {level}")
        if level == 1 and self.ch > 1:
            p.paragraph_format.page_break_before = True
        if level == 1:
            p.add_run().add_break()
        p.add_run(text.strip().replace("`", "").replace("**", ""))

    def caption(self, kind, text):
        style = "รูปที่" if kind == "fig" else "ตารางที่"
        word = "รูปที่" if kind == "fig" else "ตารางที่"
        p = self.doc.add_paragraph(style=style)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if kind == "tab":
            p.paragraph_format.keep_with_next = True
        n = self.fig_no if kind == "fig" else self.tab_no
        run = p.add_run(f"{word} {self.ch}."); run.italic = False; set_cs_size(run, 32)
        field_runs(p, f" SEQ {word} \\* ARABIC \\s 1 ", str(n), size=32)
        for r in p.runs: r.italic = False
        tail = p.add_run(" "); set_cs_size(tail, 32); tail.italic = False
        before = len(p.runs)
        self.inline(p, text.replace("`", ""), size=32)
        for r in p.runs:
            r.italic = False
            rpr = r._r.get_or_add_rPr(); ics = OxmlElement("w:iCs"); ics.set(qn("w:val"), "0"); rpr.append(ics)

    def figure(self, caption_text, src, label):
        if src.startswith("diagram:"):
            path = os.path.join(HERE, "diagrams", src.split(":", 1)[1] + ".png")
        else:
            path = os.path.join(REPO, src.replace("/", os.sep))
        p = self.doc.add_paragraph(style="Body Text")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        if os.path.exists(path):
            run = p.add_run()
            pic = run.add_picture(path, width=Cm(14.5))
            if src.startswith("diagram:"):
                max_h = Cm(19)
            else:
                max_h = Cm(10) if pic.height > pic.width * 1.6 else Cm(15)
            if pic.height > max_h:
                ratio = max_h / pic.height
                pic.height = int(max_h); pic.width = int(pic.width * ratio)
        else:
            self.missing.append(f"ch{self.ch}: image missing {src}")
            r = p.add_run(f"[รูปภาพยังไม่มี: {src}]"); r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        self.caption("fig", caption_text)

    def table(self, caption_text, rows):
        self.caption("tab", caption_text)
        ncols = max(len(r) for r in rows)
        t = self.doc.add_table(rows=len(rows), cols=ncols)
        table_borders(t)
        size = 28 if ncols <= 5 else 24 if ncols <= 8 else 20
        for i, row in enumerate(rows):
            for j in range(ncols):
                cell = t.cell(i, j)
                cp = cell.paragraphs[0]; cp.style = self.doc.styles["Table Contents"]
                txt = row[j] if j < len(row) else ""
                if i == 0:
                    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    shade(cell, "D9E2F3")
                elif txt.strip() in ("■", "□", "✓", "-", "–") or len(txt.strip()) <= 2:
                    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                self.inline(cp, txt.strip(), size=size, bold=(i == 0))
        repeat_header(t.rows[0])
        from docx.enum.table import WD_ALIGN_VERTICAL
        for c in t.rows[0].cells:
            c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        n = len(t.rows)
        for i, r in enumerate(t.rows):
            r._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            if (n <= 12 and i < n - 1) or (n > 12 and (i < 2 or i == n - 2)):
                for c in r.cells:
                    for cp in c.paragraphs:
                        cp.paragraph_format.keep_with_next = True
        if ncols > 8:  # Gantt-style: narrow month columns
            t.autofit = False
            widths = [Cm(4.8)] + [Cm(9.8 / (ncols - 1))] * (ncols - 1)
            for r in t.rows:
                for c, w in zip(r.cells, widths):
                    c.width = w
        for r in t.rows:
            for c in r.cells:
                for cp in c.paragraphs:
                    cp.paragraph_format.space_after = Pt(0); cp.paragraph_format.space_before = Pt(0)
        spacer = self.doc.add_paragraph(style="Body Text"); spacer.paragraph_format.space_after = Pt(6)

    # -- pre-pass: assign figure/table numbers so forward refs resolve
    def prepass(self, ch, lines):
        f = t = 0
        for line in lines:
            m = re.match(r"^(FIGURE|TABLE):.*\{#((?:fig|tab):[\w\-]+)\}", line)
            if not m:
                if line.startswith("FIGURE:"): f += 1
                if line.startswith("TABLE:"): t += 1
                continue
            if m.group(1) == "FIGURE":
                f += 1; self.labels[m.group(2)] = f"{ch}.{f}"
            else:
                t += 1; self.labels[m.group(2)] = f"{ch}.{t}"

    def register_refs(self, refs):
        self.local_map = {}
        for n, text in sorted(refs.items()):
            k = ref_key(text)
            if k not in self.ref_index:
                self.ref_order.append(text); self.ref_index[k] = len(self.ref_order)
            self.local_map[n] = self.ref_index[k]

    def chapter(self, lines):
        i = 0
        while i < len(lines):
            line = lines[i].rstrip()
            s = line.strip()
            if not s:
                i += 1; continue
            if s.startswith("#### "):
                p = self.para(s[5:], indent_tab=False); [setattr(r, "bold", True) for r in p.runs]
            elif s.startswith("### "): self.heading(3, s[4:])
            elif s.startswith("## "): self.heading(2, s[3:])
            elif s.startswith("# "): self.heading(1, s[2:])
            elif s.startswith("TABLE:"):
                cap = s[6:].strip(); rows = []
                i += 1
                while i < len(lines) and not lines[i].strip(): i += 1
                while i < len(lines) and lines[i].strip().startswith("|"):
                    r = lines[i].strip().strip("|")
                    cells = [c.strip() for c in re.split(r"(?<!\\)\|", r)]
                    if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                        rows.append([c.replace("\\|", "|") for c in cells])
                    i += 1
                self.tab_no += 1
                self.table(re.sub(r"\s*\{#[^}]+\}", "", cap), rows)
                continue
            elif s.startswith("FIGURE:"):
                body = s[7:].strip()
                cap, _, src = body.rpartition("|")
                label = re.search(r"\{#([^}]+)\}", cap)
                self.fig_no += 1
                self.figure(re.sub(r"\s*\{#[^}]+\}", "", cap).strip(), src.strip(), label)
            elif re.match(r"^[-*] ", s): self.list_item("•", s[2:])
            elif re.match(r"^\d+[.)] ", s):
                m = re.match(r"^(\d+)[.)] (.*)", s); self.list_item(m.group(1) + ".", m.group(2))
            elif s.startswith("|"):
                pass  # stray table without caption: skip silently is wrong -> record
            else:
                self.para(s)
            i += 1

# ---------------------------------------------------------------- front matter
def body_children(doc):
    return list(doc.element.body)


def text_of(el):
    return "".join(t.text or "" for t in el.iter(qn("w:t")))


def make_field_paragraph(doc, style_name, instr, placeholder):
    p = doc.add_paragraph(style=style_name)
    field_runs(p, instr, placeholder)
    el = p._p; el.getparent().remove(el)
    return el


def rebuild_front(doc, cover, abbrevs):
    body = doc.element.body
    kids = body_children(doc)
    # 1) cover: kids[0] logo; kids[1..22] text; kids[23] date + sectPr
    dotted_tpl = copy.deepcopy(kids[6])
    for el in kids[1:23]:
        body.remove(el)
    date_p = kids[23]
    def line(text, size=36, bold=True, space_after=0):
        el = copy.deepcopy(dotted_tpl)
        for r in el.findall(qn("w:r"))[1:]:
            el.remove(r)
        r = el.find(qn("w:r"))
        r.find(qn("w:t")).text = text
        r.find(qn("w:t")).set(qn("xml:space"), "preserve")
        rpr = r.find(qn("w:rPr"))
        if rpr is not None:
            for tag in ("w:sz", "w:szCs"):
                e = rpr.find(qn(tag))
                if e is not None: e.set(qn("w:val"), str(size))
        if "[[" in text:
            el.remove(r)
            for seg in re.split(r"(\[\[[^\]]+\]\])", text):
                if not seg: continue
                nr = copy.deepcopy(r); nr.find(qn("w:t")).text = seg.strip("[]") if seg.startswith("[[") else seg
                if seg.startswith("[["):
                    hl = OxmlElement("w:highlight"); hl.set(qn("w:val"), "yellow"); nr.find(qn("w:rPr")).append(hl)
                el.append(nr)
        date_p.addprevious(el)
    for item in cover["lines"]:
        line(item if isinstance(item, str) else item[0], size=(36 if isinstance(item, str) else item[1]))
    # date
    runs = date_p.findall(qn("w:r"))
    for r in runs[1:]:
        date_p.remove(r)
    runs[0].find(qn("w:t")).text = cover["date"].strip("[]")
    if cover["date"].startswith("[["):
        hl = OxmlElement("w:highlight"); hl.set(qn("w:val"), "yellow"); runs[0].find(qn("w:rPr")).append(hl)

    # 2) TOC / figure index / table index: clear between heading and next page-break paragraph
    def clear_after(container, heading_text, style_name, instr, placeholder):
        kids = list(container)
        start = next(i for i, el in enumerate(kids) if text_of(el).strip() == heading_text and el.tag == qn("w:p"))
        j = start + 1
        while j < len(kids) and kids[j].tag in (qn("w:p"), qn("w:bookmarkStart"), qn("w:bookmarkEnd")) \
                and not kids[j].findall(".//" + qn("w:br")):
            j += 1
        for el in kids[start + 1:j]:
            container.remove(el)
        kids[start].addnext(make_field_paragraph(doc, style_name, instr, placeholder))
    sdt_content = next(el for el in body if el.tag == qn("w:sdt")).find(qn("w:sdtContent"))
    clear_after(sdt_content, "สารบัญ", "toc 1", ' TOC \\o "1-3" \\h \\z \\u ', "(คลิกขวา > Update Field เพื่อสร้างสารบัญ)")
    clear_after(body, "รายการรูปภาพ", "Figure Index 1", ' TOC \\h \\z \\c "รูปที่" ', "(Update Field)")
    clear_after(body, "รายการตาราง", "Figure Index 1", ' TOC \\h \\z \\c "ตารางที่" ', "(Update Field)")

    # 3) abbreviation table
    kids = body_children(doc)
    tbl = next(el for el in kids if el.tag == qn("w:tbl"))
    rows = tbl.findall(qn("w:tr"))
    tpl_row = copy.deepcopy(rows[0])
    for r in rows:
        tbl.remove(r)
    for a in sorted(abbrevs, key=lambda s: s.upper()):
        nr = copy.deepcopy(tpl_row)
        cells = nr.findall(qn("w:tc"))
        for c, val in zip(cells, (a, abbrevs[a])):
            t = c.find(".//" + qn("w:t")); t.text = val
        tbl.append(nr)

    # 4) drop template chapters: keep through the thaiLetters section paragraph
    kids = body_children(doc)
    end = next(i for i, el in enumerate(kids) if el.tag == qn("w:p") and el.find(".//" + qn("w:pgNumType")) is not None
               and el.find(".//" + qn("w:pgNumType")).get(qn("w:fmt")) == "thaiLetters")
    final_sect = kids[-1]
    for el in kids[end + 1:-1]:
        body.remove(el)
    pg = OxmlElement("w:pgNumType"); pg.set(qn("w:start"), "1")
    final_sect.find(qn("w:cols")).addprevious(pg)


def main(out):
    base = os.path.join(HERE, "base.docx")
    dotx_to_docx(DOTX, base)
    doc = Document(base)
    cover = json.load(open(os.path.join(HERE, "cover.json"), encoding="utf8"))

    chapters, all_abbr, all_diagrams = [], {}, {}
    for files in CHAPTER_FILES:
        parts = []
        for f in files:
            p = os.path.join(HERE, os.environ.get("CHAPTERS", "chapters"), f)
            if os.path.exists(p):
                parts.append(split_chapter(open(p, encoding="utf8").read()))
        chapters.append(parts)
        for _, d, a, _r in parts:
            all_diagrams.update(d)
            for k, v in a.items(): all_abbr.setdefault(k, v)
    json.dump(all_diagrams, open(os.path.join(HERE, "diagrams", "specs.json"), "w", encoding="utf8"), ensure_ascii=False, indent=1)

    rebuild_front(doc, cover, all_abbr)
    b = Builder(doc)
    for ch, parts in enumerate(chapters, 1):
        b.prepass(ch, [l for lines, *_ in parts for l in lines])
    for ch, parts in enumerate(chapters, 1):
        b.ch = ch; b.fig_no = b.tab_no = 0
        for lines, _d, _a, refs in parts:
            b.register_refs(refs)
            b.chapter(lines)
    # bibliography
    p = doc.add_paragraph(style="Title"); p.paragraph_format.page_break_before = True; p.add_run("บรรณานุกรม")
    for n, text in enumerate(b.ref_order, 1):
        q = doc.add_paragraph(style="Body Text")
        q.paragraph_format.left_indent = Cm(1.2); q.paragraph_format.first_line_indent = Cm(-1.2)
        q.paragraph_format.tab_stops.add_tab_stop(Cm(1.2))
        q.alignment = WD_ALIGN_PARAGRAPH.LEFT
        q.add_run(f"[{n}]\t{text.replace('`', '')}")
    # ask Word to refresh fields on open
    settings = doc.settings.element
    uf = OxmlElement("w:updateFields"); uf.set(qn("w:val"), "true"); settings.append(uf)
    doc.core_properties.title = cover.get("doc_title", "")
    doc.save(out)
    print("saved", out, "| refs", len(b.ref_order), "| abbr", len(all_abbr), "| diagrams", len(all_diagrams))
    for m in b.missing: print("WARN", m)


if __name__ == "__main__":
    main(sys.argv[1])
